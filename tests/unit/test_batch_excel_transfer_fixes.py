from io import BytesIO
from types import SimpleNamespace
from unittest import IsolatedAsyncioTestCase, TestCase
from unittest.mock import AsyncMock, Mock, patch
from zipfile import ZipFile

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.api.v1.routers.tasks import router
from src.application.dto import BatchFilters
from src.application.exceptions import BatchCsvFileError, BatchExportFileError
from src.application.importers.batch_excel import BatchExcelParser
from src.application.reports.batch_csv import BatchCsvExporter
from src.application.reports.batch_excel_export import (
    EXCEL_CELL_MAX_CHARACTERS,
    BatchExcelExporter,
)
from src.application.services.batch_export_service import BatchExportService
from tests.unit.test_batch_details_cache import batch
from tests.unit.test_batch_transfer import excel_bytes


class ExcelExportTests(TestCase):
    def test_boundary_text_is_preserved_and_longer_text_is_rejected(self):
        row = batch(is_closed=False)
        row.task_description = "Ж" * EXCEL_CELL_MAX_CHARACTERS
        data = BatchExcelExporter().generate([row])
        result = BatchExcelParser().parse(data)
        self.assertEqual(
            result.valid_rows[0].data.task_description, row.task_description
        )
        self.assertFalse(result.errors)
        row.task_description = "X" * 40000
        with self.assertRaisesRegex(
            BatchExportFileError, "task_description.*40000.*32767"
        ):
            BatchExcelExporter().generate([row])
        self.assertEqual(len(row.task_description), 40000)

    def test_worksheet_traversal_grows_linearly_and_formula_like_text_stays_text(self):
        count = 300
        rows = []
        for index in range(count):
            row = batch(is_closed=False)
            row.task_description = f"=Task-{index}"
            rows.append(row)
        traversed = 0

        class CountingCells(dict):
            def __iter__(self):
                nonlocal traversed
                traversed += len(self)
                return super().__iter__()

        original_init = Worksheet.__init__

        def initialize(sheet, *args, **kwargs):
            original_init(sheet, *args, **kwargs)
            sheet._cells = CountingCells()

        with patch.object(Worksheet, "__init__", initialize):
            data = BatchExcelExporter().generate(rows)
        self.assertLess(traversed, 50 * count * 15)
        workbook = load_workbook(BytesIO(data))
        self.addCleanup(workbook.close)
        sheet = workbook.active
        self.assertEqual(sheet.max_row, count + 1)
        self.assertEqual(sheet["B2"].value, "=Task-0")
        self.assertEqual(sheet["B2"].data_type, "s")
        self.assertEqual(sheet.cell(count + 1, 2).value, "=Task-299")


class ExportServiceTests(IsolatedAsyncioTestCase):
    async def test_rejected_workbook_is_not_uploaded(self):
        row = batch(is_closed=False)
        row.task_description = "X" * 40000
        uow = SimpleNamespace(
            batches=SimpleNamespace(list_for_export=AsyncMock(return_value=[row])),
            rollback=AsyncMock(),
        )
        storage = SimpleNamespace(upload=Mock())
        service = BatchExportService(uow, BatchCsvExporter(), storage)
        with self.assertRaises(BatchExportFileError):
            await service.export_and_upload(filters=BatchFilters(), format="excel")
        storage.upload.assert_not_called()


class ExcelFileErrorTests(TestCase):
    def test_malformed_sheet_and_workbook_xml_are_public_file_errors(self):
        original = excel_bytes()
        for damaged_name in ["xl/worksheets/sheet1.xml", "xl/workbook.xml"]:
            with self.subTest(file=damaged_name):
                output = BytesIO()
                with (
                    ZipFile(BytesIO(original)) as source,
                    ZipFile(output, "w") as target,
                ):
                    for name in source.namelist():
                        content = source.read(name)
                        if name == damaged_name:
                            content = content[:-3]
                        target.writestr(name, content)
                with self.assertRaisesRegex(
                    BatchCsvFileError, "Invalid or damaged Excel"
                ):
                    BatchExcelParser().parse(output.getvalue())

    def test_status_exposes_expected_file_errors_but_hides_internal_errors(self):
        app = FastAPI()
        app.include_router(router, prefix="/api/v1")
        for error, expected in [
            (
                BatchExportFileError(
                    "Excel field exceeds the 32767-character cell limit"
                ),
                "Excel field exceeds the 32767-character cell limit",
            ),
            (
                BatchCsvFileError("Invalid or damaged Excel file"),
                "Invalid or damaged Excel file",
            ),
            (
                RuntimeError("private database details"),
                "Task failed; check worker logs",
            ),
        ]:
            with self.subTest(error=type(error).__name__):
                task = SimpleNamespace(
                    id="task-1",
                    backend=SimpleNamespace(
                        get_task_meta=Mock(
                            return_value={"status": "FAILURE", "result": error}
                        )
                    ),
                )
                with (
                    patch(
                        "src.tasks.task_status.celery_app.AsyncResult",
                        return_value=task,
                    ),
                    TestClient(app) as client,
                ):
                    response = client.get("/api/v1/tasks/task-1")
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.json()["error"], expected)
