from io import BytesIO
from smtplib import SMTPException
from types import SimpleNamespace
from unittest import IsolatedAsyncioTestCase, TestCase
from unittest.mock import AsyncMock, Mock, patch
from uuid import uuid4

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.api.dependencies.batches import get_batch_service
from src.api.exception_handlers import register_exception_handlers
from src.api.v1.routers.batches import router
from src.application.dto.stored_report import StoredReport
from src.application.services.batch_report_file_service import BatchReportFileService
from src.application.services.batch_report_service import BatchReportService
from src.application.services.email_service import EmailService
from src.core.config import settings
from src.domain.exceptions.batch import BatchNotFoundError
from src.tasks.report_tasks import generate_batch_report
from tests.unit.test_batch_details_cache import batch


class BatchReportFileTests(IsolatedAsyncioTestCase):
    async def test_both_formats_and_read_transaction_released_before_upload(self):
        uow = SimpleNamespace(
            batches=SimpleNamespace(
                get_with_products=AsyncMock(return_value=batch(is_closed=False))
            ),
            rollback=AsyncMock(),
        )
        storage = SimpleNamespace(upload=Mock())

        def upload(*args):
            uow.rollback.assert_awaited()

        storage.upload.side_effect = upload
        service = BatchReportFileService(BatchReportService(uow), storage, "UTC")
        for format, mime, extension in [
            (
                "excel",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "xlsx",
            ),
            ("pdf", "application/pdf", "pdf"),
        ]:
            result = await service.generate_and_upload(42, format)
            args = storage.upload.call_args.args
            self.assertEqual(args[0], "reports")
            self.assertTrue(result.object_name.endswith("." + extension))
            self.assertEqual(args[3], mime)
            if format == "excel":
                workbook = load_workbook(BytesIO(args[2]))
                self.assertEqual(len(workbook.sheetnames), 3)
                workbook.close()
            else:
                self.assertTrue(args[2].startswith(b"%PDF"))
        storage.upload.side_effect = RuntimeError("unavailable")
        with self.assertRaises(RuntimeError):
            await service.generate_and_upload(42, "excel")


class BatchReportTaskTests(TestCase):
    def test_notification_queue_failure_does_not_lose_report(self):
        report = StoredReport(uuid4(), "reports", "example.xlsx")
        for error in [None, OSError("broker down")]:
            with (
                self.subTest(error=error),
                patch(
                    "src.tasks.report_tasks._generate_batch_report",
                    new=AsyncMock(return_value=report),
                ),
                patch(
                    "src.tasks.report_tasks.send_report_notification.delay",
                    return_value=SimpleNamespace(id="email-task"),
                    side_effect=error,
                ),
            ):
                result = generate_batch_report.run(42, "excel", "person@example.com")
                self.assertEqual(result["report_id"], str(report.report_id))
                self.assertEqual(
                    result["email_status"], "queue_failed" if error else "queued"
                )

    def test_no_notification_without_email_and_no_notification_after_generation_error(
        self,
    ):
        report = StoredReport(uuid4(), "reports", "example.xlsx")
        with (
            patch(
                "src.tasks.report_tasks._generate_batch_report",
                new=AsyncMock(return_value=report),
            ),
            patch("src.tasks.report_tasks.send_report_notification.delay") as send,
        ):
            self.assertEqual(
                generate_batch_report.run(42)["email_status"], "not_requested"
            )
            send.assert_not_called()
        with (
            patch(
                "src.tasks.report_tasks._generate_batch_report",
                new=AsyncMock(side_effect=RuntimeError("upload failed")),
            ),
            patch("src.tasks.report_tasks.send_report_notification.delay") as send,
        ):
            with self.assertRaises(RuntimeError):
                generate_batch_report.run(42, "excel", "person@example.com")
            send.assert_not_called()

    def test_smtp_message_link_tls_and_failure(self):
        config = settings.model_copy(
            update={
                "smtp_security": "starttls",
                "smtp_username": "user",
                "smtp_password": None,
            }
        )
        report_id = uuid4()
        with patch("src.application.services.email_service.smtplib.SMTP") as smtp:
            connection = smtp.return_value.__enter__.return_value
            EmailService(config).send_report_notification(
                "person@example.com", 42, report_id, "pdf"
            )
            connection.starttls.assert_called_once()
            message = connection.send_message.call_args.args[0]
            self.assertIn(str(report_id), message.get_content())
            self.assertIn("format=pdf", message.get_content())
            self.assertEqual(message["To"], "person@example.com")
            connection.send_message.side_effect = SMTPException("failed")
            with self.assertRaises(SMTPException):
                EmailService(config).send_report_notification(
                    "person@example.com", 42, report_id, "pdf"
                )


class BatchReportApiTests(TestCase):
    def test_start_validation_missing_batch_and_download(self):
        app = FastAPI()
        register_exception_handlers(app)
        app.include_router(router, prefix="/api/v1")
        service = SimpleNamespace(
            get_by_id=AsyncMock(return_value=batch(is_closed=False))
        )
        app.dependency_overrides[get_batch_service] = lambda: service
        with (
            TestClient(app) as client,
            patch(
                "src.api.v1.routers.batches.generate_batch_report_task.delay",
                return_value=SimpleNamespace(id="task-1"),
            ) as delay,
        ):
            response = client.post(
                "/api/v1/batches/42/reports",
                json={"format": "pdf", "email": "person@example.com"},
            )
            self.assertEqual(response.status_code, 202, response.text)
            delay.assert_called_once_with(
                batch_id=42, format="pdf", email="person@example.com"
            )
            delay.reset_mock()
            for data in [{"format": "bad"}, {"email": "bad"}]:
                self.assertEqual(
                    client.post("/api/v1/batches/42/reports", json=data).status_code,
                    422,
                )
            delay.assert_not_called()
            service.get_by_id.side_effect = BatchNotFoundError(42)
            self.assertEqual(
                client.post("/api/v1/batches/42/reports", json={}).status_code, 404
            )
            delay.assert_not_called()
            with patch(
                "src.application.services.report_download_service.ReportDownloadService.get_batch_report_url",
                new=AsyncMock(return_value="http://example.com/report"),
            ) as download:
                report_id = uuid4()
                response = client.get(
                    f"/api/v1/batches/42/reports/{report_id}/download?format=pdf",
                    follow_redirects=False,
                )
                self.assertEqual(response.status_code, 307)
                download.assert_awaited_once_with(42, report_id, "pdf")
