from datetime import UTC, date, datetime
from io import BytesIO
from types import SimpleNamespace
from unittest import IsolatedAsyncioTestCase, TestCase
from unittest.mock import AsyncMock, Mock, patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy.exc import DataError, OperationalError

from src.api.v1.routers.batches import router
from src.api.v1.schemas.batch_transfer import BatchExportRequest
from src.application.dto import BatchFilters
from src.application.exceptions import BatchCsvFileError
from src.application.importers.batch_csv import REQUIRED_HEADERS, BatchCsvParser
from src.application.importers.batch_excel import BatchExcelParser
from src.application.reports.batch_csv import BatchCsvExporter
from src.application.services.batch_export_service import BatchExportService
from src.application.services.batch_import_service import BatchImportService
from src.storage.batch_list_cache import BatchListCache
from tests.unit.test_batch_details_cache import batch


def excel_bytes(rows=None):
    workbook = Workbook()
    sheet = workbook.active
    headers = sorted(REQUIRED_HEADERS)
    sheet.append(headers)
    defaults = {
        "task_description": "Task",
        "work_center_id": 1,
        "shift": "day",
        "team": "A",
        "batch_number": 111,
        "batch_date": date(2026, 10, 2),
        "nomenclature": "Bolt",
        "ekn_code": "EKN",
        "shift_start": datetime(2026, 10, 2, 8, tzinfo=UTC).replace(tzinfo=None),
        "shift_end": datetime(2026, 10, 2, 16, tzinfo=UTC).replace(tzinfo=None),
    }
    for row in rows or [{}]:
        data = defaults | row
        sheet.append([data[key] for key in headers])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class BatchTransferParserTests(TestCase):
    def test_dates_numbers_formulas_and_row_errors(self):
        result = BatchExcelParser().parse(
            excel_bytes(
                [
                    {},
                    {"batch_number": 3000000000},
                    {"task_description": "=1+1"},
                    {"batch_number": 112},
                ]
            )
        )
        self.assertEqual([r.row_number for r in result.valid_rows], [2, 5])
        self.assertEqual([r.row_number for r in result.errors], [3, 4])
        self.assertEqual(
            result.valid_rows[0].data.shift_start, datetime(2026, 10, 2, 3, tzinfo=UTC)
        )
        self.assertIn("formulas", result.errors[1].reason)

    def test_broken_and_oversized_excel_are_file_errors(self):
        with self.assertRaises(BatchCsvFileError):
            BatchExcelParser().parse(b"not excel")
        with (
            patch(
                "src.application.importers.batch_excel.settings.batch_import_max_uncompressed_bytes",
                1,
            ),
            self.assertRaises(BatchCsvFileError),
        ):
            BatchExcelParser().parse(excel_bytes())
        with (
            patch(
                "src.application.importers.batch_excel.settings.batch_import_max_rows",
                1,
            ),
            self.assertRaises(BatchCsvFileError),
        ):
            BatchExcelParser().parse(excel_bytes([{}, {}]))

    def test_export_request_validates_range(self):
        for filters in [
            {"date_from": "2026-10-02", "date_to": "2026-10-01"},
            {"batch_date": "2026-10-02", "date_from": "2026-10-01"},
        ]:
            with self.assertRaises(ValueError):
                BatchExportRequest.model_validate({"filters": filters})


class BatchTransferServiceTests(IsolatedAsyncioTestCase):
    async def test_row_data_error_rolls_back_and_continues_but_system_error_stops(self):
        parser = BatchExcelParser()
        service = SimpleNamespace(
            create=AsyncMock(
                side_effect=[
                    SimpleNamespace(id=1),
                    DataError("SQL", {}, Exception("private")),
                    SimpleNamespace(id=3),
                ]
            ),
            uow=SimpleNamespace(rollback=AsyncMock()),
        )
        importer = BatchImportService(parser, service)
        result = await importer.import_file(
            excel_bytes([{}, {"batch_number": 112}, {"batch_number": 113}])
        )
        self.assertEqual([r.batch_id for r in result.processed], [1, 3])
        self.assertEqual(result.failed[0].row_number, 3)
        self.assertNotIn("private", result.failed[0].reason)
        service.uow.rollback.assert_awaited_once()
        service.create.side_effect = OperationalError("SQL", {}, Exception("offline"))
        with self.assertRaises(OperationalError):
            await importer.import_file(excel_bytes())

    async def test_export_formats_use_exports_bucket_and_can_be_imported(self):
        row = batch(is_closed=False)
        row.task_description = "=example"
        uow = SimpleNamespace(
            batches=SimpleNamespace(list_for_export=AsyncMock(return_value=[row])),
            rollback=AsyncMock(),
        )
        storage = SimpleNamespace(upload=Mock())
        service = BatchExportService(uow, BatchCsvExporter(), storage)
        for format in ["csv", "excel"]:
            result = await service.export_and_upload(
                filters=BatchFilters(), format=format
            )
            args = storage.upload.call_args.args
            self.assertEqual(args[0], "exports")
            self.assertEqual(result.bucket, "exports")
            if format == "excel":
                parsed = BatchExcelParser().parse(args[2])
                workbook = load_workbook(BytesIO(args[2]))
                self.assertEqual(workbook.active["B2"].data_type, "s")
                workbook.close()
            else:
                parsed = BatchCsvParser().parse(args[2])
            self.assertEqual(len(parsed.valid_rows), 1)
            self.assertFalse(parsed.errors)

    async def test_cache_keys_separate_ranges(self):
        cache = BatchListCache(AsyncMock(get=AsyncMock(return_value="v1")), 60)
        first = await cache.make_key(
            filters=BatchFilters(date_from=date(2026, 10, 1)), offset=0, limit=20
        )
        second = await cache.make_key(
            filters=BatchFilters(date_from=date(2026, 10, 2)), offset=0, limit=20
        )
        self.assertNotEqual(first, second)


class BatchTransferApiTests(TestCase):
    def test_import_and_export_parameters_and_validation(self):
        app = FastAPI()
        app.include_router(router, prefix="/api/v1")
        with (
            TestClient(app) as client,
            patch(
                "src.api.v1.routers.batches.export_batches_file_task.delay",
                return_value=SimpleNamespace(id="export-task"),
            ) as export,
        ):
            response = client.post(
                "/api/v1/batches/export",
                json={
                    "format": "excel",
                    "filters": {"date_from": "2026-10-01", "date_to": "2026-10-02"},
                },
            )
            self.assertEqual(response.status_code, 202, response.text)
            self.assertEqual(
                export.call_args.kwargs["filters"]["date_from"], "2026-10-01"
            )
            self.assertEqual(export.call_args.kwargs["format"], "excel")
            self.assertEqual(
                client.post(
                    "/api/v1/batches/export",
                    json={
                        "filters": {"date_from": "2026-10-02", "date_to": "2026-10-01"}
                    },
                ).status_code,
                422,
            )
            with (
                patch(
                    "src.api.v1.routers.batches.create_object_storage",
                    return_value=SimpleNamespace(upload=Mock()),
                ) as storage,
                patch(
                    "src.api.v1.routers.batches.import_batches_file_task.delay",
                    return_value=SimpleNamespace(id="import-task"),
                ) as task,
            ):
                response = client.post(
                    "/api/v1/batches/import",
                    files={"file": ("batches.xlsx", excel_bytes())},
                )
                self.assertEqual(response.status_code, 202, response.text)
                self.assertEqual(task.call_args.kwargs["format"], "excel")
                self.assertEqual(
                    storage.return_value.upload.call_args.args[0], "imports"
                )
                self.assertEqual(
                    client.post(
                        "/api/v1/batches/import", files={"file": ("file.exe", b"bad")}
                    ).status_code,
                    422,
                )
            self.assertEqual(
                client.get(
                    "/api/v1/batches?date_from=2026-10-02&date_to=2026-10-01"
                ).status_code,
                422,
            )
