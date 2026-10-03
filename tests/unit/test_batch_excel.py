from datetime import UTC, datetime
from io import BytesIO
from types import SimpleNamespace
from unittest import IsolatedAsyncioTestCase
from unittest.mock import AsyncMock, patch

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.application.reports.batch_excel import BatchExcelGenerator
from src.application.services.batch_report_service import BatchReportService
from tests.unit.test_batch_details_cache import batch


class BatchExcelTests(IsolatedAsyncioTestCase):
    async def test_three_sheets_products_statistics_timezone_and_literal_codes(self):
        row = batch(is_closed=False)
        now = datetime(2026, 10, 2, 8, 0, tzinfo=UTC)
        row.products = [
            SimpleNamespace(
                id=1,
                unique_code="=1+1",
                is_aggregated=True,
                aggregated_at=now,
                created_at=now,
            ),
            SimpleNamespace(
                id=2,
                unique_code="P2",
                is_aggregated=False,
                aggregated_at=None,
                created_at=now,
            ),
        ]
        service = BatchReportService(
            SimpleNamespace(
                batches=SimpleNamespace(get_with_products=AsyncMock(return_value=row))
            )
        )
        data = await service.collect(42)
        workbook = load_workbook(
            BytesIO(BatchExcelGenerator("Asia/Yekaterinburg").generate(data))
        )
        self.assertEqual(
            workbook.sheetnames, ["Информация о партии", "Продукция", "Статистика"]
        )
        sheet = workbook["Продукция"]
        self.assertEqual(sheet["B2"].value, "=1+1")
        self.assertEqual(sheet["B2"].data_type, "s")
        self.assertEqual(sheet["D2"].value, "2026-10-02 13:00:00+05:00")
        self.assertIsNone(sheet["D3"].value)
        stats = workbook["Статистика"]
        self.assertEqual([stats[f"B{i}"].value for i in range(2, 6)], [2, 1, 1, 0.5])
        self.assertEqual(stats["B5"].number_format, "0.00%")
        workbook.close()

    async def test_large_report_has_linear_cell_traversals_and_preserves_rows(self):
        count = 1000
        row = batch(is_closed=False)
        now = datetime(2026, 10, 2, 8, tzinfo=UTC)
        row.products = [
            SimpleNamespace(
                id=index,
                unique_code="=1+1" if index == 1 else f"P-{index}",
                is_aggregated=False,
                aggregated_at=None,
                created_at=now,
            )
            for index in range(1, count + 1)
        ]
        service = BatchReportService(
            SimpleNamespace(
                batches=SimpleNamespace(get_with_products=AsyncMock(return_value=row))
            )
        )
        report = await service.collect(42)
        traversed = 0

        class CountingCells(dict):
            def __iter__(self):
                nonlocal traversed
                traversed += len(self)
                return super().__iter__()

        original_init = Worksheet.__init__

        def initialize(sheet, *args, **kwargs):
            original_init(sheet, *args, **kwargs)
            sheet._cells = CountingCells()

        with patch.object(Worksheet, "__init__", initialize):
            data = BatchExcelGenerator("UTC").generate(report)
        # Bound total work, without depending on machine speed or wall-clock timing.
        self.assertLess(traversed, 50 * count * 5)
        workbook = load_workbook(BytesIO(data))
        self.addCleanup(workbook.close)
        products = workbook["Продукция"]
        self.assertEqual(products.max_row, count + 1)
        self.assertEqual(products["B2"].value, "=1+1")
        self.assertEqual(products["B2"].data_type, "s")
        self.assertEqual(products.cell(count + 1, 2).value, f"P-{count}")

    async def test_empty_batch_still_generates_valid_workbook(self):
        service = BatchReportService(
            SimpleNamespace(
                batches=SimpleNamespace(
                    get_with_products=AsyncMock(return_value=batch(is_closed=False))
                )
            )
        )
        workbook = load_workbook(
            BytesIO(BatchExcelGenerator("UTC").generate(await service.collect(42)))
        )
        self.assertEqual(workbook["Продукция"].max_row, 1)
        self.assertEqual(
            [workbook["Статистика"][f"B{i}"].value for i in range(2, 6)], [0, 0, 0, 0]
        )
        workbook.close()
