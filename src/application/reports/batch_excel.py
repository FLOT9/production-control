from datetime import date, datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.application.dto.batch_report import BatchReportData


class BatchExcelGenerator:
    def __init__(self, timezone: str) -> None:
        self.timezone = ZoneInfo(timezone)

    def _time(self, value: datetime | None) -> str:
        return (
            value.astimezone(self.timezone).isoformat(sep=" ", timespec="seconds")
            if value
            else ""
        )

    @staticmethod
    def _append(sheet: Worksheet, values: list[str | int | float | date]) -> None:
        cells = []
        for column, value in enumerate(values, start=1):
            # append() assigns the final row without querying worksheet dimensions.
            cell = Cell(sheet, row=1, column=column, value=value)
            # Keep codes starting with '=' as text, without scanning previous rows.
            if isinstance(value, str):
                cell.data_type = "s"
            cells.append(cell)
        sheet.append(cells)

    def generate(self, report: BatchReportData) -> bytes:
        workbook = Workbook()
        info = workbook.active
        if info is None:
            raise RuntimeError("Workbook has no active sheet")
        info.title = "Информация о партии"
        products = workbook.create_sheet("Продукция")
        stats = workbook.create_sheet("Статистика")
        batch = report.batch

        self._append(info, ["Поле", "Значение"])
        info_rows: list[tuple[str, str | int | float | date]] = [
            ("ID партии", batch.id),
            ("Номер партии", batch.batch_number),
            ("Дата партии", batch.batch_date),
            ("Задание", batch.task_description),
            ("ID рабочего центра", batch.work_center_id),
            ("Смена", batch.shift),
            ("Бригада", batch.team),
            ("Номенклатура", batch.nomenclature),
            ("Код ЕКН", batch.ekn_code),
            ("Начало смены", self._time(batch.shift_start)),
            ("Окончание смены", self._time(batch.shift_end)),
            ("Статус партии", "Закрыта" if batch.is_closed else "Открыта"),
            ("Время закрытия", self._time(batch.closed_at)),
            ("Создана", self._time(batch.created_at)),
            ("Обновлена", self._time(batch.updated_at)),
            ("Отчёт сформирован", self._time(report.generated_at)),
            ("Часовой пояс", self.timezone.key),
        ]
        for title, value in info_rows:
            self._append(info, [title, value])

        self._append(
            products, ["ID", "Код продукта", "Агрегирован", "Время агрегации", "Создан"]
        )
        for product in report.products:
            self._append(
                products,
                [
                    product.id,
                    product.unique_code,
                    "Да" if product.is_aggregated else "Нет",
                    self._time(product.aggregated_at),
                    self._time(product.created_at),
                ],
            )
        if report.products:
            products.auto_filter.ref = products.dimensions

        self._append(stats, ["Показатель", "Значение"])
        for title, value in [
            ("Всего продуктов", report.statistics.total_products),
            ("Агрегировано", report.statistics.aggregated_products),
            ("Осталось", report.statistics.pending_products),
            ("Процент агрегации", report.statistics.aggregation_percent / 100),
        ]:
            self._append(stats, [title, value])
        stats["B5"].number_format = "0.00%"

        for sheet in workbook.worksheets:
            sheet.freeze_panes = "A2"
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(fill_type="solid", fgColor="244062")
            for index, column in enumerate(sheet.columns, start=1):
                width = max(len(str(cell.value or "")) for cell in column) + 2
                sheet.column_dimensions[get_column_letter(index)].width = min(width, 65)

        output = BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()
