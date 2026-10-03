from collections.abc import Sequence
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from src.application.exceptions import BatchExportFileError
from src.application.reports.batch_csv import HEADERS
from src.data.models import Batch

EXCEL_CELL_MAX_CHARACTERS = 32767


class BatchExcelExporter:
    def generate(self, batches: Sequence[Batch]) -> bytes:
        workbook = Workbook()
        try:
            sheet = workbook.active
            if sheet is None:
                raise RuntimeError("Workbook has no worksheet")
            sheet.title = "Партии"
            sheet.append(list(HEADERS))
            for batch in batches:
                cells = []
                for column, header in enumerate(HEADERS, start=1):
                    value = getattr(batch, header)
                    if hasattr(value, "isoformat"):
                        value = value.isoformat()
                    if (
                        isinstance(value, str)
                        and len(value) > EXCEL_CELL_MAX_CHARACTERS
                    ):
                        raise BatchExportFileError(
                            f"Excel export cannot store field '{header}' of batch {batch.id}: "
                            f"{len(value)} characters exceeds the "
                            f"{EXCEL_CELL_MAX_CHARACTERS}-character cell limit"
                        )
                    # append() assigns the final row. Prepare types before adding cells,
                    # avoiding a full worksheet scan to find the last row.
                    export_cell = Cell(sheet, row=1, column=column, value=value)
                    if isinstance(value, str):
                        export_cell.data_type = "s"
                    cells.append(export_cell)
                sheet.append(cells)
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for index, cell in enumerate(sheet[1], start=1):
                sheet.column_dimensions[get_column_letter(index)].width = min(
                    len(str(cell.value)) + 8, 40
                )
            data = BytesIO()
            workbook.save(data)
            return data.getvalue()
        finally:
            workbook.close()
