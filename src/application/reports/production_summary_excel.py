from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from src.application.dto import ProductionSummaryRow

HEADERS = (
    "Код участка",
    "Название участка",
    "Номер партии",
    "Дата партии",
    "Смена",
    "Бригада",
    "Номенклатура",
    "Код ЕКН",
    "Статус партии",
    "Всего продуктов",
    "Агрегировано продуктов",
)

DANGEROUS_EXCEL_PREFIXES = ("=", "+", "-", "@")
HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
DATE_FORMAT = "dd.mm.yyyy"
COLUMN_WIDTHS = {
    "A": 16,
    "B": 28,
    "C": 14,
    "D": 14,
    "E": 14,
    "F": 20,
    "G": 28,
    "H": 16,
    "I": 16,
    "J": 18,
    "K": 24,
}


def _escape_excel_formula(value: str) -> str:
    if value.startswith(DANGEROUS_EXCEL_PREFIXES):
        return f"'{value}"

    return value


def _format_worksheet(worksheet: Worksheet) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 30

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT

    for column, width in COLUMN_WIDTHS.items():
        worksheet.column_dimensions[column].width = width

    for cell in worksheet["D"][1:]:
        cell.number_format = DATE_FORMAT


class ProductionSummaryExcelGenerator:
    def generate(self, rows: list[ProductionSummaryRow]) -> bytes:
        workbook = Workbook()
        worksheet = workbook.active
        if not isinstance(worksheet, Worksheet):
            raise TypeError("Workbook has no active worksheet")
        worksheet.title = "Производственная сводка"
        worksheet.append(HEADERS)

        for row in rows:
            worksheet.append(
                (
                    _escape_excel_formula(row.work_center_identifier),
                    _escape_excel_formula(row.work_center_name),
                    row.batch_number,
                    row.batch_date,
                    _escape_excel_formula(row.shift),
                    _escape_excel_formula(row.team),
                    _escape_excel_formula(row.nomenclature),
                    _escape_excel_formula(row.ekn_code),
                    "Закрыта" if row.is_closed else "Открыта",
                    row.products_total,
                    row.products_aggregated,
                )
            )

        _format_worksheet(worksheet)

        buffer = BytesIO()
        workbook.save(buffer)

        return buffer.getvalue()
