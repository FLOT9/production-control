from io import BytesIO

# Only the exception type is imported; openpyxl handles XML parsing.
from xml.etree.ElementTree import ParseError  # nosec B405
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError

from src.application.dto.batch_import import (
    BatchCsvParseResult,
    BatchImportRow,
    BatchImportRowError,
    ParsedBatchImportRow,
)
from src.application.exceptions import BatchCsvFileError
from src.application.importers.batch_csv import REQUIRED_HEADERS
from src.core.config import settings


class BatchExcelParser:
    def parse(self, data: bytes) -> BatchCsvParseResult:
        try:
            with ZipFile(BytesIO(data)) as archive:
                if (
                    sum(item.file_size for item in archive.infolist())
                    > settings.batch_import_max_uncompressed_bytes
                ):
                    raise BatchCsvFileError("Excel unpacked size exceeds the limit")
            workbook = load_workbook(
                BytesIO(data), read_only=True, data_only=False, keep_links=False
            )
            try:
                sheet = workbook.worksheets[0]
                sheet.reset_dimensions()
                rows = sheet.iter_rows()
                header_cells = next(rows, ())
                headers = [
                    str(cell.value).strip() if cell.value is not None else ""
                    for cell in header_cells
                ]
                missing = REQUIRED_HEADERS - set(headers)
                if missing:
                    raise BatchCsvFileError(
                        "Excel is missing required headers: "
                        + ", ".join(sorted(missing))
                    )
                nonempty_headers = [header for header in headers if header]
                if len(nonempty_headers) != len(set(nonempty_headers)):
                    raise BatchCsvFileError("Excel contains duplicate headers")
                valid = []
                errors = []
                for number, cells in enumerate(rows, start=2):
                    if number > settings.batch_import_max_rows + 1:
                        raise BatchCsvFileError("Excel row limit exceeded")
                    if all(cell.value is None for cell in cells):
                        continue
                    if any(cell.data_type == "f" for cell in cells):
                        errors.append(
                            BatchImportRowError(
                                number, "Excel formulas are not allowed"
                            )
                        )
                        continue
                    if any(cell.value is not None for cell in cells[len(headers) :]):
                        errors.append(
                            BatchImportRowError(
                                number, "Excel row has more values than headers"
                            )
                        )
                        continue
                    row_data = {
                        header: cells[i].value if i < len(cells) else None
                        for i, header in enumerate(headers)
                        if header
                    }
                    try:
                        row = BatchImportRow.model_validate(row_data)
                    except ValidationError as error:
                        errors.append(BatchImportRowError(number, str(error)))
                    else:
                        valid.append(ParsedBatchImportRow(number, row))
                return BatchCsvParseResult(valid, errors)
            finally:
                workbook.close()
        except (
            BadZipFile,
            InvalidFileException,
            KeyError,
            ValueError,
            IndexError,
            ParseError,
        ) as error:
            raise BatchCsvFileError("Invalid or damaged Excel file") from error
